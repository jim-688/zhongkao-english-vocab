#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
中考英语核心词表构建管线 v2（本地优先版）
数据源: words_20tian.json(词集+释义) + sguo4.xlsx(音标) + freq_50k.txt(词频)
输出: JSON + Excel，高频在前、同频段乱序
"""
import csv, json, random, re, os, sys

ROOT = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.environ.get("VOCAB_DATA_DIR", os.path.join(ROOT, "data"))
OUTPUT_DIR = os.environ.get("VOCAB_OUTPUT_DIR", ROOT)
WORDS_SRC = os.path.join(DATA_DIR, "words_20tian.json")
SGUO4 = os.path.join(DATA_DIR, "sguo4.xlsx")
FREQ = os.path.join(DATA_DIR, "freq_50k.txt")
OUT_JSON = os.path.join(OUTPUT_DIR, "zhongkao_vocab.json")
OUT_XLSX = os.path.join(OUTPUT_DIR, "zhongkao_vocab.xlsx")
SEED = 20260814

def load_wordlist():
    return json.load(open(WORDS_SRC, encoding="utf-8"))

def load_prev_phonetic():
    """复用已生成的 zhongkao_vocab.json 中的音标（API/有道补齐的结果），防止重建覆盖。"""
    if os.path.exists(OUT_JSON):
        try:
            prev = json.load(open(OUT_JSON, encoding="utf-8"))
            return {r["word"]: r.get("phonetic", "") for r in prev}
        except Exception:
            return {}
    return {}

def load_sguo4(path):
    """读取 sguo4.xlsx: 单词\t英音\t美音\t释义（多行记录：释义可能跨行）"""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    idx = {}
    cur = None
    for row in ws.iter_rows(values_only=True):
        cell0 = row[0]
        if cell0 and str(cell0).strip():
            word = str(cell0).strip()
            ph = ""
            if row[1] and str(row[1]).strip():
                ph = str(row[1]).strip()
            if row[2] and str(row[2]).strip() and str(row[2]).strip() != ph:
                ph += " " + str(row[2]).strip()
            cur = {"phonetic": ph, "meaning": ""}
            idx[word] = cur
            if row[3] and str(row[3]).strip():
                cur["meaning"] = str(row[3]).strip()
        elif cur and row[3] and str(row[3]).strip():
            m = str(row[3]).strip()
            cur["meaning"] = cur["meaning"] + "；" + m if cur["meaning"] else m
    wb.close()
    return idx

def load_freq(path):
    """freq_50k.txt: 每行一个词或 'word freq'？先探测格式"""
    freq = {}
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            parts = line.split()
            if len(parts) >= 2 and parts[1].replace(".", "").isdigit():
                freq[parts[0].lower()] = float(parts[1])
            else:
                freq[line.lower()] = len(freq)
    return freq

def parse_pos_from_meaning(meaning):
    parts = meaning.split("；")[0]
    tags = re.findall(r"\b(n|vt|vi|adj|adv|prep|conj|pron|num|int|aux|det|art|modalv|abbr)\.", parts)
    return " ".join(tags) if tags else ""

POS_TAGS = ("n", "vt", "vi", "adj", "adv", "prep", "conj", "pron", "num", "int", "aux", "det", "art", "abbr")

def simplify_meaning(meaning, max_pos=3):
    """零基础版：每个词性只保留第一个义项，最多 max_pos 个词性。"""
    parts = [p.strip() for p in meaning.split("；") if p.strip()]
    result = []
    seen = set()
    for part in parts:
        m = re.match(r"^(" + "|".join(POS_TAGS) + r")\.?\s*(.*)$", part)
        if m and m.group(1) not in seen:
            seen.add(m.group(1))
            result.append(f"{m.group(1)}. {m.group(2)}")
            if len(seen) >= max_pos:
                break
    if not result:
        return meaning[:60]
    return "；".join(result)

def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    words = load_wordlist()
    print(f"词集: {len(words)} 词")
    sguo = load_sguo4(SGUO4)
    prev_ph = load_prev_phonetic()
    print(f"sguo4 音标库: {len(sguo)} 词")
    freq = load_freq(FREQ)
    print(f"词频库: {len(freq)} 词")

    rows = []
    for w in words:
        word = w["word"]
        s = sguo.get(word, {})
        phonetic = prev_ph.get(word) or s.get("phonetic", "")
        rows.append({
            "word": word,
            "phonetic": phonetic,
            "pos": parse_pos_from_meaning(w["meaning"]),
            "meaning": w["meaning"],
            "meaning_simple": simplify_meaning(w["meaning"]),
            "frq": freq.get(word.lower()),
            "from_sguo": bool(s),
        })

    missing_phon = [r["word"] for r in rows if not r["phonetic"]]
    print(f"缺音标: {len(missing_phon)} 词 ({missing_phon[:15]})")
    missing_frq = [r["word"] for r in rows if r["frq"] is None]
    print(f"缺词频: {len(missing_frq)} 词 (前15: {missing_frq[:15]})")

    # 排序: frq 高(数值大=高频?) 先用缺失的垫底，有值的按值降序(假定高值=高频)，段内 seed 乱序
    def frq_key(r):
        return - (r["frq"] or 0)
    rows.sort(key=frq_key)
    rng = random.Random(SEED)
    for start in range(0, len(rows), 500):
        seg = rows[start:start + 500]
        rng.shuffle(seg)
        rows[start:start + 500] = seg

    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=1)
    print(f"JSON -> {OUT_JSON}")

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "中考核心词"
    ws.append(["序号", "单词", "音标", "词性", "中文释义"])
    for i, r in enumerate(rows, 1):
        ws.append([i, r["word"], r["phonetic"], r["pos"], r["meaning"]])
    wb.save(OUT_XLSX)
    print(f"Excel -> {OUT_XLSX}")

    print("\n前10词(高频):")
    for r in rows[:10]:
        print(f"  {r['word']} {r['phonetic']} [{r['pos']}] {r['meaning'][:40]}")

if __name__ == "__main__":
    main()
